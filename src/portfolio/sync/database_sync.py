"""Money3本地同步组件；从既有账本适配器迁入并裁剪非同步入口。"""

from __future__ import annotations
from collections import defaultdict, deque


from datetime import datetime, timedelta, timezone


from decimal import Decimal


import hashlib


import json


import math


from pathlib import Path


import re


import sqlite3


import pandas as pd


from portfolio.sync.order_id_utils import normalize_order_id


from portfolio.paths import PROJECT_DIR as BASE_DIR

PARSER_VERSION = "database-sync-2026-10-20.2"


SCALE = Decimal(100_000_000)


START = "2026-07-01"


TRADE_COLUMNS = ["订单ID", "交易时间", "Ticker", "方向", "成交数量", "成交均价", "成交金额", "谁"]


TRADE_SHEETS = {"LONGBRIDGE": "longbridge", "BBAE": "bbae", "FIRSTRADE": "firstrade", "SCHWAB": "schwab"}


class TradeFlowWorkbook:
    """券商同步仅修改流水；数据库导入必须重新读取已保存的文件。"""

    def __init__(self, base):
        self.path = base / "交易流水.xlsx"
        if not self.path.is_file():
            raise FileNotFoundError(f"交易流水是成交入库来源，不能省略：{self.path}")
        self.digest = hashlib.sha256(self.path.read_bytes()).hexdigest()
        self.frames = {}
        self.changed = set()
        self.cash_frames = {}
        self.cash_changed = set()
        self.cancelled_changed = False
        with pd.ExcelFile(self.path, engine="openpyxl") as excel:
            self.sheet_names = set(excel.sheet_names)
            for broker, sheet in TRADE_SHEETS.items():
                if sheet in excel.sheet_names:
                    frame = excel.parse(sheet, dtype={"订单ID": str})
                    required = set(TRADE_COLUMNS) - {"订单ID"}
                    if not required.issubset(frame.columns):
                        raise ValueError(
                            f"交易流水.xlsx [{sheet}] 缺少字段：{sorted(required - set(frame.columns))}"
                        )
                    self.frames[broker] = frame
                else:
                    self.frames[broker] = pd.DataFrame(columns=TRADE_COLUMNS)
            self.cancelled = (
                excel.parse("schwab_撤销记录", dtype={"订单ID": str})
                if "schwab_撤销记录" in excel.sheet_names
                else pd.DataFrame(columns=TRADE_COLUMNS)
            )
            for sheet in ("firstrade_资金来源", "schwab_资金"):
                if sheet in excel.sheet_names:
                    self.cash_frames[sheet] = excel.parse(sheet, dtype={"活动ID": str})
            if "内部划转" in excel.sheet_names:
                self.transfer_frame = excel.parse("内部划转")
            else:
                self.transfer_frame = None

    def update_cash(self, sheet, frame):
        old = self.cash_frames.get(sheet)
        if sheet == "firstrade_资金来源" and old is not None:
            dates = pd.to_datetime(frame["日期"], errors="coerce").dropna()
            if dates.empty:
                return  # 无日期的券商导出不能声明撤销既有资金活动。
            old_dates = pd.to_datetime(old["日期"], errors="coerce")
            retained = old[~old_dates.between(dates.min(), dates.max())]
            parts = [part for part in (retained, frame) if not part.empty]
            frame = pd.concat(parts, ignore_index=True) if parts else frame
        if old is not None and _json(old.to_dict("records")) == _json(frame.to_dict("records")):
            return
        self.cash_frames[sheet] = frame.copy()
        self.cash_changed.add(sheet)

    def read(self, broker):
        return self.frames[broker].copy()

    def assignments(self, broker):
        frame = self.read(broker)
        if broker == "SCHWAB" and not self.cancelled.empty:
            frame = pd.concat([self.cancelled, frame], ignore_index=True)
        result = {}
        for row in frame.to_dict("records"):
            identity = normalize_order_id(row.get("订单ID"))
            if identity:
                result[identity] = _text(row.get("谁"))
        return result

    def update(self, broker, frame, *, replace=True):
        from portfolio.sync.sync_and_split_operations import _same_transaction_rows

        old = self.read(broker)
        columns = (
            TRADE_COLUMNS if broker in {"LONGBRIDGE", "SCHWAB"} or "订单ID" in old else TRADE_COLUMNS[1:]
        )
        new = frame.reindex(columns=columns).copy()
        if not replace:
            ids = set(new.get("订单ID", pd.Series(dtype=str)).map(normalize_order_id))
            retained = old[
                ~old.get("订单ID", pd.Series("", index=old.index)).map(normalize_order_id).isin(ids)
            ]
            parts = [part for part in (retained.reindex(columns=columns), new) if not part.empty]
            new = pd.concat(parts, ignore_index=True) if parts else new
        if not new.empty:
            new["交易时间"] = pd.to_datetime(new["交易时间"], errors="raise")
            new = new.sort_values("交易时间", ascending=False, kind="stable").reset_index(drop=True)
        if _same_transaction_rows(old, new):
            return 0
        if broker == "SCHWAB":
            active_ids = set(new["订单ID"].map(normalize_order_id))
            removed = old[~old["订单ID"].map(normalize_order_id).isin(active_ids)]
            parts = [part for part in (self.cancelled, removed) if not part.empty]
            archived = (pd.concat(parts, ignore_index=True) if parts else self.cancelled.copy()).reindex(
                columns=TRADE_COLUMNS
            )
            archived = archived[~archived["订单ID"].map(normalize_order_id).isin(active_ids)]
            self.cancelled = archived.drop_duplicates(subset=["订单ID"], keep="last")
            self.cancelled_changed = True
        self.frames[broker] = new
        self.changed.add(broker)
        return len(new)

    def save(self):
        import os
        from uuid import uuid4
        from openpyxl import load_workbook
        from portfolio.sync.sync_and_split_operations import _write_df_to_sheet

        if hashlib.sha256(self.path.read_bytes()).hexdigest() != self.digest:
            raise ValueError("交易流水.xlsx 在同步期间被修改，请保存并关闭文件后重试；未覆盖你的修改。")
        if not self.changed and not self.cancelled_changed and not self.cash_changed:
            return
        temporary = self.path.with_name(f".{self.path.stem}.{uuid4().hex}.xlsx")
        workbook = load_workbook(self.path)
        try:
            for broker in sorted(self.changed):
                _write_df_to_sheet(workbook, TRADE_SHEETS[broker], self.frames[broker])
            if self.cancelled_changed:
                _write_df_to_sheet(workbook, "schwab_撤销记录", self.cancelled)
            for sheet in sorted(self.cash_changed):
                # 资金活动 ID 同样按文本保存，避免 Excel 截断长整数。
                frame = self.cash_frames[sheet].copy()
                if "活动ID" in frame:
                    frame["活动ID"] = frame["活动ID"].map(_text)
                frame = frame.astype(object).where(pd.notna(frame), None)
                _write_df_to_sheet(workbook, sheet, frame)
            workbook.save(temporary)
            if hashlib.sha256(self.path.read_bytes()).hexdigest() != self.digest:
                raise ValueError("保存前发现交易流水.xlsx 被修改，已停止覆盖，请重试。")
            os.replace(temporary, self.path)
        finally:
            workbook.close()
            temporary.unlink(missing_ok=True)
        sheets = [TRADE_SHEETS[b] for b in sorted(self.changed)] + sorted(self.cash_changed)
        print(f"[SYNC] 券商数据已先保存到交易流水.xlsx：{', '.join(sheets)}。")


def _now():
    return datetime.now(timezone.utc).isoformat(timespec="microseconds").replace("+00:00", "Z")


def _text(value):
    return "" if value is None or pd.isna(value) else str(value).strip()


def _scaled(value):
    from portfolio.sync.import_excel_database import fixed_e8

    result = fixed_e8(value, allow_round=True)
    return 0 if result is None else result


def _value(value):
    if value is None:
        return None
    if isinstance(value, dict):
        return {str(key): _value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_value(item) for item in value]
    if isinstance(value, (datetime, pd.Timestamp)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if pd.isna(value):
        return None
    if hasattr(value, "item"):
        return _value(value.item())
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _json(value):
    return json.dumps(
        _value(value), ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False
    )


def _sha(value):
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _insert(connection, table, **values):
    names = list(values)
    return connection.execute(
        f"INSERT INTO {table} ({','.join(names)}) VALUES ({','.join('?' for _ in names)})",
        [values[name] for name in names],
    ).lastrowid


def _batch(
    connection,
    *,
    scope,
    source_kind,
    records,
    source_path=None,
    digest=None,
    sheet="API",
    request_start=None,
    request_end=None,
    cursor=None,
):
    """归档本次实际读取的原始行。成功状态仅在整个事务校验完成后设置。"""
    content = digest or _sha(_json(records))
    batch_key = f"{scope}:{content}:{PARSER_VERSION}"
    if source_kind == "API":
        batch_key += ":" + (cursor or _now())
    # 内容 A→B→A 也属于一次新导入，不能因历史中曾出现过 A 而丢掉本次来源顺序。
    if connection.execute("SELECT 1 FROM import_batches WHERE batch_key=?", (batch_key,)).fetchone():
        batch_key += ":" + _now()
    batch_id = _insert(
        connection,
        "import_batches",
        batch_key=batch_key,
        source_kind=source_kind,
        source_scope=scope,
        source_path=source_path,
        content_sha256=content,
        parser_version=PARSER_VERSION,
        request_start=request_start,
        request_end=request_end,
        resume_cursor=cursor,
        status="PENDING",
        imported_at_utc=_now(),
        record_count=len(records),
        input_role="LIVE_SOURCE",
    )
    ids = {}
    for ordinal, record in enumerate(records, 2):
        entity = _text(
            record.get("activityId") or record.get("活动ID") or record.get("订单ID") or record.get("order_id")
        )
        payload = _json({"fields": record})
        ids[ordinal] = _insert(
            connection,
            "raw_records",
            import_batch_id=batch_id,
            source_locator=f"{sheet}:row:{ordinal}",
            source_entity_id=entity or None,
            original_time_text=_text(record.get("交易时间") or record.get("time") or record.get("成交时间")),
            payload_json=payload,
            payload_sha256=_sha(payload),
            parse_status="ACCEPTED",
            parse_message="数据库同步原始记录；非成交行保留，不伪装成买卖",
            source_sheet=sheet,
            source_row=ordinal,
            record_kind="DATA",
        )
    return batch_id, ids


def _account(connection, broker):
    rows = connection.execute(
        "SELECT id FROM broker_accounts WHERE broker_code=? AND enabled=1", (broker,)
    ).fetchall()
    if len(rows) > 1:
        raise ValueError(f"{broker} 存在多个启用账户，当前同步适配器需明确选择账户，不能混合合并。")
    if rows:
        return rows[0]["id"]
    local = broker in {"BBAE", "LONGBRIDGE"}
    return _insert(
        connection,
        "broker_accounts",
        broker_code=broker,
        account_key="LEGACY_DEFAULT",
        display_name=f"{broker}（本地账本账户）",
        base_currency="USD",
        book_timezone="Asia/Shanghai" if local else "America/New_York",
        day_cutoff="08:00" if local else "00:00",
    )


def _instrument(connection, symbol, source_id):
    from portfolio.sync.sync_and_split_operations import _ticker, is_ignored_ticker

    symbol = _ticker(symbol)
    if is_ignored_ticker(symbol):
        raise ValueError("同步成交存在不受支持的美股证券代码，请先核对市场。")
    row = connection.execute(
        "SELECT id FROM instruments WHERE market='US' AND symbol=?", (symbol,)
    ).fetchone()
    if row:
        return row["id"]
    identifier = _insert(
        connection,
        "instruments",
        market="US",
        symbol=symbol,
        currency="USD",
        asset_type="UNKNOWN",
        category="待确认",
        source_record_id=source_id,
    )
    _insert(
        connection,
        "instrument_aliases",
        provider="DATABASE_SYNC",
        market="US",
        source_symbol=symbol,
        instrument_id=identifier,
        source_record_id=source_id,
    )
    return identifier


def _skip_unsupported_security(connection, source_id, raw_symbol):
    """沿用美股主流程已确认的排除规则，原行仍完整归档。"""
    from portfolio.sync.sync_and_split_operations import is_ignored_ticker

    symbol = _text(raw_symbol).upper()
    if symbol in {"", "NAN", "NONE", "NAT"}:
        raise ValueError("已成交来源缺少有效证券代码，本次数据库同步未提交。")
    if symbol and is_ignored_ticker(symbol):
        connection.execute(
            "UPDATE raw_records SET parse_status='IGNORED',parse_message=? WHERE id=?",
            ("既有美股主流程排除的港股或数字证券代码；原始行已保留，未混入美股成交", source_id),
        )
        return True
    return False


def _read_trades(connection, broker):
    rows = connection.execute(
        "SELECT t.*,i.symbol FROM broker_trades t JOIN broker_accounts a ON a.id=t.account_id "
        "JOIN instruments i ON i.id=t.instrument_id WHERE a.broker_code=? "
        "AND t.record_state='ACTIVE' AND t.accounting_role='ACCOUNT_LEDGER' AND i.market='US' ORDER BY t.id",
        (broker,),
    ).fetchall()
    records = [
        {
            "_db_id": row["id"],
            "订单ID": row["broker_order_id"] or "",
            "交易时间": pd.Timestamp(row["book_time"]),
            "Ticker": row["symbol"],
            "方向": "买入" if row["side"] == "BUY" else "卖出",
            "成交数量": float(Decimal(row["quantity_e8"]) / SCALE),
            "成交均价": float(Decimal(row["price_e8"] or 0) / SCALE),
            "成交金额": float(Decimal(row["calculation_amount_e8"]) / SCALE),
            "谁": row["assignment_text"] or "",
        }
        for row in rows
    ]
    return pd.DataFrame(records, columns=[*TRADE_COLUMNS, "_db_id"])


def _assignments(connection, broker):
    # 撤销后再次恢复的活动也保留原人工归属。
    rows = connection.execute(
        "SELECT t.broker_order_id,t.assignment_text FROM broker_trades t "
        "JOIN broker_accounts a ON a.id=t.account_id WHERE a.broker_code=? "
        "AND t.broker_order_id IS NOT NULL AND t.accounting_role='ACCOUNT_LEDGER'",
        (broker,),
    ).fetchall()
    result = {}
    for row in rows:
        if row["broker_order_id"] in result:
            raise ValueError(f"{broker} 数据库存在重复订单 ID，无法安全同步。")
        result[row["broker_order_id"]] = row["assignment_text"] or ""
    return result


def _signature(row):
    return (
        pd.Timestamp(row["交易时间"]).strftime("%Y-%m-%d %H:%M:%S"),
        str(row["Ticker"]).upper(),
        _text(row["方向"]),
        _scaled(abs(float(row["成交数量"]))),
        _scaled(row["成交均价"]),
        _scaled(abs(float(row["成交金额"]))),
    )


def _apply_trades(
    connection, broker, frame, *, source_ids, reverse_missing=False, authoritative_source=False
):
    """按 ID 或带出现序号的完整成交特征匹配，保留相同数值的独立成交。"""
    account = _account(connection, broker)
    old = connection.execute(
        "SELECT t.* FROM broker_trades t JOIN instruments i ON i.id=t.instrument_id "
        "WHERE t.account_id=? AND t.accounting_role='ACCOUNT_LEDGER' AND i.market='US'",
        (account,),
    ).fetchall()
    if len([row for row in old if row["broker_order_id"]]) != len(
        {row["broker_order_id"] for row in old if row["broker_order_id"]}
    ):
        raise ValueError(f"{broker} 数据库存在重复订单 ID，无法安全同步。")
    by_id = {row["id"]: row for row in old}
    by_order = {row["broker_order_id"]: row for row in old if row["broker_order_id"]}
    by_signature = defaultdict(deque)
    for row in old:
        if row["record_state"] != "ACTIVE":
            continue
        symbol = connection.execute(
            "SELECT symbol FROM instruments WHERE id=?", (row["instrument_id"],)
        ).fetchone()[0]
        by_signature[
            (
                row["book_time"],
                symbol,
                "买入" if row["side"] == "BUY" else "卖出",
                row["quantity_e8"],
                row["price_e8"] or 0,
                row["calculation_amount_e8"],
            )
        ].append(row)
    used, changed = set(), 0
    for _, item in frame.iterrows():
        signature = _signature(item)
        order_id = normalize_order_id(item.get("订单ID")) or ""
        original = None
        if pd.notna(item.get("_db_id")):
            original = by_id.get(int(item["_db_id"]))
        if original is None and order_id:
            original = by_order.get(order_id)
        if original is None and not order_id:
            queue = by_signature[signature]
            while queue and queue[0]["id"] in used:
                queue.popleft()
            original = queue.popleft() if queue else None
        if original is not None and original["id"] in used:
            raise ValueError(f"{broker} 新来源多次引用同一成交，已停止数据库更新。")
        if signature[2] not in {"买入", "卖出"} or signature[3] <= 0 or signature[4] < 0:
            raise ValueError(f"{broker} 成交方向、数量或价格无效。")
        raw_row = item.get("__source_row")
        raw_id = source_ids.get(int(raw_row)) if pd.notna(raw_row) else None
        if original is None and authoritative_source and raw_id is not None and not order_id:
            original = next(
                (row for row in old if row["source_record_id"] == raw_id and row["id"] not in used), None
            )
        if raw_id is None:
            if original is None:
                raise ValueError(f"{broker} 新成交缺少原始记录来源。")
            raw_id = original["source_record_id"]
        security = _instrument(connection, signature[1], raw_id)
        values = dict(
            account_id=account,
            instrument_id=security,
            broker_order_id=order_id or None,
            record_granularity="ORDER_TOTAL" if order_id else "SOURCE_RECORD",
            side="BUY" if signature[2] == "买入" else "SELL",
            business_date=signature[0][:10],
            book_time=signature[0],
            quantity_e8=signature[3],
            price_e8=signature[4],
            calculation_amount_e8=signature[5],
            amount_basis="NET" if broker == "SCHWAB" else "LEGACY",
            currency="USD",
            record_state="ACTIVE",
            accounting_role="ACCOUNT_LEDGER",
            assignment_text=_text(item.get("谁")),
        )
        if broker == "SCHWAB":
            values["net_amount_e8"] = signature[5] * (1 if values["side"] == "SELL" else -1)
        if original is None:
            identity = f"{broker}:ID:{order_id}" if order_id else f"{broker}:RAW:{raw_id}"
            trade_id = _insert(
                connection, "broker_trades", trade_key=_sha(identity), source_record_id=raw_id, **values
            )
            used.add(trade_id)
            changed += 1
        else:
            trade_id = original["id"]
            used.add(trade_id)
            if any(original[key] != value for key, value in values.items()):
                values["source_record_id"] = raw_id
                keys = list(values)
                connection.execute(
                    f"UPDATE broker_trades SET {','.join(key + '=?' for key in keys)} WHERE id=?",
                    [values[key] for key in keys] + [trade_id],
                )
                changed += 1
            elif authoritative_source and original["source_record_id"] != raw_id:
                connection.execute(
                    "UPDATE broker_trades SET source_record_id=? WHERE id=?", (raw_id, trade_id)
                )
    if reverse_missing:
        for original in old:
            if (
                original["record_state"] == "ACTIVE"
                and original["id"] not in used
                and (authoritative_source or original["business_date"] >= START)
            ):
                connection.execute(
                    "UPDATE broker_trades SET record_state='REVERSED' WHERE id=?", (original["id"],)
                )
                changed += 1
    return changed


def _source_changed(connection, path, base):
    relative = path.relative_to(base).as_posix()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    last = connection.execute(
        "SELECT content_sha256 FROM import_batches WHERE source_path=? "
        "AND status='SUCCESS' ORDER BY id DESC LIMIT 1",
        (relative,),
    ).fetchone()
    return relative, digest, last is None or last[0] != digest


def _import_bbae(connection, base, ledger):
    from portfolio.sync.sync_and_split_operations import (
        _number,
        _ticker,
        _direction,
        _beijing_bookkeeping_time,
        _to_eastern_naive,
        _reconcile_bbae_orders,
    )

    path = base / "原始数据/手动维护/交易订单数据bbae.xlsx"
    if not path.exists():
        return
    relative, digest, changed = _source_changed(connection, path, base)
    if not changed:
        print("[SKIP] BBAE 源文件未变化，保留交易流水中的成交及人工归属。")
        return
    raw = pd.read_excel(path, sheet_name="股票")
    needed = {"订单状态", "成交数量", "成交价", "成交时间", "股票代码", "方向"}
    if not needed.issubset(raw.columns):
        raise ValueError(f"BBAE 源文件缺少字段：{sorted(needed - set(raw.columns))}")
    _, ids = _batch(
        connection,
        scope="BBAE:ORDER_FILE",
        source_kind="EXCEL",
        source_path=relative,
        digest=digest,
        records=raw.to_dict("records"),
        sheet="股票",
    )
    rows, ignored = [], 0
    for index, item in raw.iterrows():
        if _text(item["订单状态"]) not in {"已成", "已成交"}:
            continue
        if _skip_unsupported_security(connection, ids[index + 2], item["股票代码"]):
            ignored += 1
            continue
        quantity, price = abs(_number(item["成交数量"])), _number(item["成交价"])
        raw_time = pd.to_datetime(item["成交时间"], errors="coerce")
        if quantity <= 0 or price <= 0 or pd.isna(raw_time):
            raise ValueError(f"BBAE 已成交来源第 {index + 2} 行数量、价格或时间无效，未推进导入状态。")
        if raw_time < pd.Timestamp(START):
            continue
        rows.append(
            {
                "交易时间": _beijing_bookkeeping_time(raw_time),
                "Ticker": _ticker(item["股票代码"]),
                "方向": _direction(item["方向"]),
                "成交数量": quantity,
                "成交均价": price,
                "成交金额": round(quantity * price, 2),
                "谁": "",
                "__source_row": index + 2,
                "__原始北京时间": raw_time,
                "__旧版美东时间": _to_eastern_naive(raw_time),
            }
        )
    if rows:
        merged, _, ambiguous, unresolved = _reconcile_bbae_orders(ledger.read("BBAE"), pd.DataFrame(rows))
        if ambiguous or unresolved:
            raise ValueError("BBAE 新文件存在无法唯一匹配的成交或归属；本次数据库同步已回滚，请核对后重试。")
        count = ledger.update("BBAE", merged)
    else:
        count = 0
    if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
        raise ValueError("BBAE 源文件在读取过程中发生变化，请关闭正在修改该文件的程序后重试。")
    print(f"[SYNC] BBAE 新文件已核对，待保存流水 {count} 笔。")
    if ignored:
        print(f"[SKIP] BBAE 已归档并排除 {ignored} 条既有美股流程不支持的港股或数字证券记录。")


def _firstrade_cash(connection, raw, account, ids, relative):
    from portfolio.sync.sync_and_split_operations import _number

    times = pd.to_datetime(raw["日期"], errors="coerce").dropna()
    low, high = (
        (times.min().strftime("%Y-%m-%d"), times.max().strftime("%Y-%m-%d")) if not times.empty else ("", "")
    )
    existing = connection.execute(
        "SELECT e.* FROM cash_events e JOIN raw_records r ON r.id=e.source_record_id "
        "JOIN import_batches b ON b.id=r.import_batch_id WHERE e.account_id=? "
        "AND e.record_state='ACTIVE' AND ((b.source_path IN (?,?) AND e.business_date BETWEEN ? AND ?) "
        "OR b.source_scope='TRADE_FLOW:CASH:FIRSTRADE' OR r.source_sheet='firstrade_资金来源')",
        (
            account,
            str(Path(relative).with_suffix(".csv")),
            str(Path(relative).with_suffix(".xlsx")),
            low,
            high,
        ),
    ).fetchall()
    candidates = defaultdict(deque)
    for row in existing:
        candidates[(row["business_date"], row["event_type"], row["amount_e8"])].append(row)
    deposits = {"入金", "存款", "deposit", "wire funds received", "wire_in"}
    withdrawals = {"出金", "取款", "提款", "withdraw", "wire funds sent", "wire_out"}
    used = set()
    for index, item in raw.iterrows():
        source_type = _text(item["交易类别"])
        day = pd.to_datetime(item["日期"], errors="coerce")
        if source_type.lower() not in deposits | withdrawals:
            continue
        try:
            valid = not pd.isna(day) and math.isfinite(float(item.get("金额")))
        except (ValueError, TypeError):
            valid = False
        if not valid:
            raise ValueError(f"交易流水.xlsx [firstrade_资金来源] 第 {index + 2} 行资金日期或金额无效。")
        kind = "DEPOSIT" if source_type.lower() in deposits else "WITHDRAWAL"
        amount = _scaled(abs(_number(item.get("金额")))) * (1 if kind == "DEPOSIT" else -1)
        day = day.strftime("%Y-%m-%d")
        queue = candidates[(day, kind, amount)]
        if queue:
            matched = queue.popleft()
            used.add(matched["id"])
            connection.execute(
                "UPDATE cash_events SET source_record_id=?,description=?,source_event_type=? WHERE id=?",
                (ids[index + 2], _text(item.get("说明")), source_type, matched["id"]),
            )
            continue
        _insert(
            connection,
            "cash_events",
            cash_event_key=f"FIRSTRADE:RAW:{ids[index + 2]}",
            account_id=account,
            event_type=kind,
            business_date=day,
            amount_e8=amount,
            currency="USD",
            accounting_role="CAPITAL",
            description=_text(item.get("说明")),
            source_event_type=source_type,
            source_record_id=ids[index + 2],
        )
    for row in existing:
        if row["id"] not in used:
            connection.execute("UPDATE cash_events SET record_state='REVERSED' WHERE id=?", (row["id"],))


def _import_firstrade(connection, base, ledger):
    from portfolio.sync.sync_and_split_operations import (
        _number,
        _ticker,
        _direction,
        _refresh_firstrade_source_window,
    )

    path = base / "原始数据/手动维护/交易订单数据firstrade.csv"
    if not path.exists():
        path = path.with_suffix(".xlsx")
    if not path.exists():
        return
    relative, digest, changed = _source_changed(connection, path, base)
    if not changed and "firstrade_资金来源" in ledger.cash_frames:
        print("[SKIP] Firstrade 源文件未变化，保留交易流水中的成交及人工归属。")
        return
    if path.suffix == ".csv":
        raw = None
        for encoding in ("utf-8-sig", "utf-8", "gbk", "gb18030", "latin1"):
            try:
                raw = pd.read_csv(path, encoding=encoding)
                break
            except UnicodeError:
                continue
        if raw is None:
            raise ValueError("无法读取 Firstrade CSV 编码。")
    else:
        raw = pd.read_excel(path)
    needed = {"日期", "数量", "价格", "交易类别", "代号"}
    if not needed.issubset(raw.columns):
        raise ValueError(f"Firstrade 源文件缺少字段：{sorted(needed - set(raw.columns))}")
    _, ids = _batch(
        connection,
        scope="FIRSTRADE:ORDER_FILE",
        source_kind="CSV" if path.suffix == ".csv" else "EXCEL",
        source_path=relative,
        digest=digest,
        records=raw.to_dict("records"),
        sheet="CSV" if path.suffix == ".csv" else "Sheet1",
    )
    rows, ignored = [], 0
    for index, item in raw.iterrows():
        when = pd.to_datetime(item["日期"], errors="coerce")
        quantity, price = abs(_number(item["数量"])), _number(item["价格"])
        direction = _direction(item["交易类别"])
        if direction not in {"买入", "卖出"}:
            continue
        if _skip_unsupported_security(connection, ids[index + 2], item["代号"]):
            ignored += 1
            continue
        if quantity <= 0 or price <= 0 or pd.isna(when):
            raise ValueError(f"Firstrade 成交来源第 {index + 2} 行数量、价格或时间无效，未推进导入状态。")
        if when < pd.Timestamp(START):
            continue
        rows.append(
            {
                "交易时间": when,
                "Ticker": _ticker(item["代号"]),
                "方向": direction,
                "成交数量": quantity,
                "成交均价": price,
                "成交金额": round(abs(_number(item.get("金额"))) or quantity * price, 2),
                "__source_row": index + 2,
            }
        )
    count = 0
    # 交易被修订成资金活动时，买卖列表可能为空；原始导出的日期范围仍是
    # 权威覆盖窗口，不能只由筛选后的买卖反推范围并留下已被删除的旧成交。
    source_dates = pd.to_datetime(raw["日期"], errors="coerce").dropna()
    source_dates = source_dates[source_dates >= pd.Timestamp(START)]
    if not source_dates.empty:
        downloaded = pd.DataFrame(rows, columns=[*TRADE_COLUMNS, "__source_row"])
        merged, *_ = _refresh_firstrade_source_window(
            ledger.read("FIRSTRADE"),
            downloaded,
            source_window=(source_dates.min(), source_dates.max()),
        )
        count = ledger.update("FIRSTRADE", merged)
    ledger.update_cash("firstrade_资金来源", raw)
    if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
        raise ValueError("Firstrade 源文件在读取过程中发生变化，请关闭正在修改该文件的程序后重试。")
    print(f"[SYNC] Firstrade 新文件已核对，待保存流水 {count} 笔；导出区间外数据已保留。")
    if ignored:
        print(f"[SKIP] Firstrade 已归档并排除 {ignored} 条既有美股流程不支持的港股或数字证券记录。")


def _sync_longbridge(connection, ledger):
    from portfolio.sync.sync_and_split_operations import (
        _longbridge_context,
        _number,
        _ticker,
        _direction,
        _beijing_bookkeeping_time,
    )
    from portfolio.sync.longbridge_credentials import CREDENTIALS_FILE

    required = bool(
        connection.execute(
            "SELECT EXISTS(SELECT 1 FROM broker_accounts WHERE broker_code='LONGBRIDGE' AND enabled=1) "
            "OR EXISTS(SELECT 1 FROM broker_trades t JOIN broker_accounts a ON a.id=t.account_id WHERE a.broker_code='LONGBRIDGE')",
        ).fetchone()[0]
    )
    context = _longbridge_context(strict=required or CREDENTIALS_FILE.is_file())
    if context is None:
        if required:
            raise RuntimeError(
                "数据库已有长桥账户或成交，但未能建立长桥交易连接，本次同步未提交。"
                "请检查凭证、Token、网络及 SDK；如只使用数据库现有账本，可添加 --skip-sync。"
            )
        print("[SKIP] 数据库尚无长桥账户及成交，未建立长桥连接；本次不导入长桥订单。")
        return
    from longbridge.openapi import OrderStatus

    start, end = datetime(2026, 7, 1), datetime.now() + timedelta(days=1)
    statuses = [OrderStatus.Filled, OrderStatus.PartialFilled]
    orders = list(context.history_orders(start_at=start, end_at=end, status=statuses) or [])
    orders += list(context.today_orders(status=statuses) or [])
    by_id = {}
    for order in orders:
        identity = normalize_order_id(getattr(order, "order_id", ""))
        if not identity:
            raise ValueError("长桥返回缺少订单 ID 的成交，无法安全去重。")
        by_id[identity] = order
    records = [
        {
            field: getattr(order, field, None)
            for field in (
                "order_id",
                "symbol",
                "side",
                "status",
                "executed_quantity",
                "executed_price",
                "price",
                "updated_at",
                "submitted_at",
            )
        }
        for order in by_id.values()
    ]
    cursor = _now()
    _, ids = _batch(
        connection,
        scope="LONGBRIDGE:ORDERS",
        source_kind="API",
        records=records,
        request_start=START,
        request_end=cursor,
        cursor=cursor,
    )
    who = ledger.assignments("LONGBRIDGE")
    rows, ignored = [], 0
    for ordinal, (identity, order) in enumerate(by_id.items(), 2):
        if _skip_unsupported_security(connection, ids[ordinal], getattr(order, "symbol", "")):
            ignored += 1
            continue
        when = _beijing_bookkeeping_time(getattr(order, "updated_at", None))
        if pd.isna(when):
            raise ValueError(f"长桥已成交订单 {identity} 缺少有效成交时间，本次数据库同步未提交。")
        if when < pd.Timestamp(START):
            continue
        quantity = _number(getattr(order, "executed_quantity", 0))
        raw_price = getattr(order, "executed_price", None)
        price = _number(getattr(order, "price", 0) if raw_price is None else raw_price)
        direction = _direction(getattr(order, "side", ""))
        if (
            not math.isfinite(quantity)
            or not math.isfinite(price)
            or quantity <= 0
            or price <= 0
            or direction not in {"买入", "卖出"}
        ):
            raise ValueError(f"长桥已成交订单 {identity} 的数量、价格或买卖方向无效，本次数据库同步未提交。")
        rows.append(
            {
                "订单ID": identity,
                "交易时间": when,
                "Ticker": _ticker(getattr(order, "symbol", "")),
                "方向": direction,
                "成交数量": quantity,
                "成交均价": price,
                "成交金额": round(quantity * price, 2),
                "谁": who.get(identity, ""),
                "__source_row": ordinal,
            }
        )
    count = ledger.update("LONGBRIDGE", pd.DataFrame(rows), replace=False) if rows else 0
    print(
        f"[{'SYNC' if count else 'SKIP'}] 长桥已核对交易流水，待保存 {count} 笔；人工归属及 API 未返回的既有成交已保留。"
    )
    if ignored:
        print(f"[SKIP] 长桥已归档并排除 {ignored} 条既有美股流程不支持的港股或数字证券记录。")


def _load_schwab_archive(connection):
    last = connection.execute(
        "SELECT id FROM import_batches WHERE source_scope='SCHWAB:TRANSACTIONS' "
        "AND status='SUCCESS' ORDER BY id DESC LIMIT 1"
    ).fetchone()
    if not last:
        return {}
    records = [
        json.loads(row[0])["fields"]
        for row in connection.execute(
            "SELECT payload_json FROM raw_records WHERE import_batch_id=? ORDER BY source_row", (last["id"],)
        )
    ]
    metadata = next((row["__metadata"] for row in records if "__metadata" in row), None)
    if metadata is None:
        raise ValueError("嘉信数据库同步档案缺少元信息，已停止，不能静默混合账户。")
    return {**metadata, "transactions": [row for row in records if "activityId" in row]}


def _corporate_actions(connection):
    return pd.read_sql_query(
        "SELECT i.symbol AS Ticker,c.effective_date AS 生效日期,c.broker_scope AS 适用渠道,"
        "c.ratio_numerator*1.0/c.ratio_denominator AS 比例 FROM corporate_actions c "
        "JOIN instruments i ON i.id=c.instrument_id WHERE c.record_state='ACTIVE'",
        connection,
    ).assign(生效日期=lambda frame: pd.to_datetime(frame["生效日期"]))


def _schwab_cash_kind(kind, amount):
    deposits = {"ACH_RECEIPT", "CASH_RECEIPT", "WIRE_IN"}
    withdrawals = {"ACH_DISBURSEMENT", "CASH_DISBURSEMENT", "WIRE_OUT"}
    if kind in deposits | withdrawals | {"ELECTRONIC_FUND"}:
        return ("DEPOSIT" if amount >= 0 else "WITHDRAWAL"), "CAPITAL"
    if kind == "DIVIDEND_OR_INTEREST":
        # 嘉信把分红和利息放在同一类型，缺少明确信息时不猜是哪一种。
        return "OTHER", "INCOME"
    return "OTHER", "INFORMATIONAL"


def _sync_schwab(connection, ledger):
    from portfolio.sync.schwab_api import get_service, normalize_transactions, live_positions
    from portfolio.sync.schwab_integration import check_quantities
    from portfolio.sync.build_investment_overviews import replay_trade_details

    service = get_service()
    if service is None:
        return
    archive = service.history(archived=_load_schwab_archive(connection), use_local_archive=False)
    new, cash = normalize_transactions(archive["transactions"])
    old = ledger.read("SCHWAB")
    if old["订单ID"].isna().any() or old["订单ID"].astype(str).str.strip().eq("").any():
        raise ValueError("交易流水.xlsx [schwab] 存在无 ID 记录，无法安全同步。")
    who = ledger.assignments("SCHWAB")
    new["谁"] = new["订单ID"].map(who).fillna("")
    known_ids = {f"SCHWAB:{item['activityId']}" for item in archive["transactions"]}
    retained = old[~old["订单ID"].isin(known_ids)]
    present = [frame for frame in (retained, new) if not frame.empty]
    desired = (pd.concat(present, ignore_index=True) if present else new.copy()).sort_values(
        "交易时间", ascending=False, kind="stable"
    )
    operations = desired.rename(columns={"成交均价": "成交均价($)", "成交金额": "成交金额($)"}).copy()
    operations["渠道"] = "SCHWAB"
    details, state = replay_trade_details(
        operations, _corporate_actions(connection), person_name="美股账户总览"
    )
    if not details.empty and (details["后持仓数量"] < -1e-6).any():
        raise ValueError("嘉信历史重放出现负数持仓，请补齐早期成交或公司行动；本次同步未写入。")
    check_quantities(
        {ticker: pos["quantity"] for (_, ticker), pos in state.items()}, live_positions(service.account())
    )
    metadata = {key: archive[key] for key in ("account_hash", "history_start", "fetched_through")}
    records = [{"__metadata": metadata}, *archive["transactions"]]
    _batch(
        connection,
        scope="SCHWAB:TRANSACTIONS",
        source_kind="API",
        records=records,
        request_start=archive["history_start"],
        request_end=archive["fetched_through"],
        cursor=archive["fetched_through"],
    )
    count = ledger.update("SCHWAB", desired)
    ledger.update_cash("schwab_资金", cash)
    print(
        f"[{'SYNC' if count else 'SKIP'}] 嘉信已核对交易流水，待保存 {count} 笔，历史股数与实盘一致；资金活动先保存到流水。"
    )


def _apply_schwab_cash(connection, cash, ids):
    """仅消费交易流水中的嘉信资金工作表，不消费 API 返回值。"""
    account = _account(connection, "SCHWAB")
    active_cash = set()
    for index, item in cash.iterrows():
        key = "SCHWAB:" + _text(item["活动ID"])
        if key in active_cash or not _text(item["活动ID"]):
            raise ValueError(f"交易流水.xlsx [schwab_资金] 第 {index + 2} 行活动 ID 缺失或重复。")
        active_cash.add(key)
        try:
            valid = not pd.isna(pd.to_datetime(item["日期"], errors="coerce")) and math.isfinite(
                float(item["净金额(USD)"])
            )
        except (ValueError, TypeError):
            valid = False
        if not valid:
            raise ValueError(f"交易流水.xlsx [schwab_资金] 第 {index + 2} 行资金日期或金额无效。")
        amount = _scaled(item["净金额(USD)"])
        kind, role = _schwab_cash_kind(item["类型"], amount)
        values = dict(
            account_id=account,
            event_type=kind,
            business_date=pd.Timestamp(item["日期"]).strftime("%Y-%m-%d"),
            amount_e8=amount,
            currency="USD",
            accounting_role=role,
            description=_text(item["说明"]),
            source_event_type=item["类型"],
            record_state="ACTIVE",
        )
        row = connection.execute("SELECT * FROM cash_events WHERE cash_event_key=?", (key,)).fetchone()
        raw_id = ids[index + 2]
        if row is None:
            _insert(connection, "cash_events", cash_event_key=key, source_record_id=raw_id, **values)
        elif any(row[field] != value for field, value in values.items()):
            fields = list(values)
            connection.execute(
                f"UPDATE cash_events SET {','.join(field + '=?' for field in fields)},source_record_id=? WHERE id=?",
                [values[field] for field in fields] + [raw_id, row["id"]],
            )
        elif row["source_record_id"] != raw_id:
            connection.execute("UPDATE cash_events SET source_record_id=? WHERE id=?", (raw_id, row["id"]))
    existing = connection.execute(
        "SELECT cash_event_key FROM cash_events WHERE account_id=? AND cash_event_key LIKE 'SCHWAB:%'",
        (account,),
    )
    for identity in {row[0] for row in existing} - active_cash:
        connection.execute(
            "UPDATE cash_events SET record_state='REVERSED' WHERE account_id=? AND cash_event_key=?",
            (account, identity),
        )


def _archive_operations(connection, identifiers, reason):
    """Keep the full replaced record and its original provenance for audit."""
    connection.execute(
        "CREATE TABLE IF NOT EXISTS person_operations_archive ("
        "id INTEGER PRIMARY KEY, original_id INTEGER NOT NULL, archived_at TEXT NOT NULL, "
        "reason TEXT NOT NULL, payload_json TEXT NOT NULL)"
    )
    has_fees = connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='person_trade_fees'"
    ).fetchone() is not None
    count = 0
    for identifier in identifiers:
        row = connection.execute("SELECT * FROM person_operations WHERE id=?", (identifier,)).fetchone()
        if row is not None:
            connection.execute(
                "INSERT INTO person_operations_archive(original_id,archived_at,reason,payload_json) VALUES(?,?,?,?)",
                (identifier, _now(), reason, _json(dict(row))),
            )
            if has_fees:
                connection.execute("DELETE FROM person_trade_fees WHERE operation_id=?", (identifier,))
            connection.execute("DELETE FROM person_operations WHERE id=?", (identifier,))
            count += 1
    return count


def _allocation_amounts(total_e8, quantities):
    """Allocate cents by largest remainder, preserving even sub-cent source totals."""
    quantity = sum(quantities)
    cents, subcent = divmod(total_e8, 1_000_000)
    amounts = [cents * q // quantity for q in quantities]
    order = sorted(range(len(quantities)), key=lambda i: (-(cents * quantities[i] % quantity), i))
    for i in order[: cents - sum(amounts)]:
        amounts[i] += 1
    amounts = [value * 1_000_000 for value in amounts]
    amounts[-1] += subcent
    return amounts


def refresh_allocations(connection):
    """All market fills follow the workbook; independent transfers remain intact."""
    people = {row["name"]: row["id"] for row in connection.execute("SELECT id,name FROM persons")}
    if "个人" not in people:
        raise ValueError("数据库缺少个人账户归属，无法拆分成交。")
    trades = connection.execute(
        "SELECT t.*,a.broker_code FROM broker_trades t JOIN broker_accounts a ON a.id=t.account_id "
        "JOIN instruments i ON i.id=t.instrument_id WHERE i.market='US' "
        "AND t.accounting_role='ACCOUNT_LEDGER' AND a.broker_code IN ('LONGBRIDGE','BBAE','FIRSTRADE','SCHWAB') "
        "ORDER BY t.book_time,t.id",
    ).fetchall()
    next_sequence = connection.execute(
        "SELECT coalesce(max(sequence_no),-1)+1 FROM person_operations"
    ).fetchone()[0]
    next_id = connection.execute("SELECT coalesce(max(id),0)+1 FROM person_operations").fetchone()[0]
    legacy = [
        row[0]
        for row in connection.execute(
            "SELECT o.id FROM person_operations o JOIN broker_accounts a ON a.id=o.account_id "
            "JOIN instruments i ON i.id=o.instrument_id WHERE o.broker_trade_id IS NULL "
            "AND o.operation_kind='MARKET_TRADE' AND i.market='US' "
            "AND a.broker_code IN ('LONGBRIDGE','BBAE','FIRSTRADE','SCHWAB')"
        )
    ]
    changed = _archive_operations(
        connection, legacy, "交易流水.xlsx为市场成交及人员分配来源，替代旧未关联市场交易"
    )
    for trade in trades:
        existing = connection.execute(
            "SELECT * FROM person_operations WHERE broker_trade_id=? "
            "AND operation_kind='MARKET_TRADE' ORDER BY sequence_no,id",
            (trade["id"],),
        ).fetchall()
        expected = []
        if trade["record_state"] == "ACTIVE":
            assignments = _text(trade["assignment_text"])
            # 保留数据库的 8 位数量精度；旧文件里的 4 位数量结果保持一致。
            # 不能先四舍五入个人余量，否则不足 0.00005 股的合法成交会消失。
            followers = [
                (name.strip(), _scaled(quantity))
                for name, quantity in re.findall(r"([^:：;；,，\s]+)\s*[:：]\s*(\d+(?:\.\d+)?)", assignments)
            ]
            residual = re.sub(r"([^:：;；,，\s]+)\s*[:：]\s*(\d+(?:\.\d+)?)", "", assignments)
            if residual.strip(" \t\r\n,，;；"):
                raise ValueError(
                    f"{trade['broker_code']} 成交 {trade['id']} 的【谁】格式无效，请填写“姓名:数量”。"
                )
            total_e8 = trade["quantity_e8"]
            assigned = sum(quantity for _, quantity in followers)
            if assigned - total_e8 >= int(SCALE):
                raise ValueError(
                    f"{trade['broker_code']} 成交 {trade['id']} 的人工分摊数量超过总成交数量，已停止。"
                )
            tolerated_excess = max(0, assigned - total_e8)
            personal = max(0, total_e8 - assigned)
            if personal:
                followers.append(("个人", personal))
            if len([name for name, _ in followers]) != len({name for name, _ in followers}):
                raise ValueError(
                    f"{trade['broker_code']} 成交 {trade['id']} 的【谁】重复填写同一人员，请合并后重试。"
                )
            followers = [(name, quantity) for name, quantity in followers if quantity > 0]
            if tolerated_excess:
                # 保留人工股数及原成交单位成本，不把已接受的小额差异暗中摊给其他人。
                allocations = [
                    _scaled(
                        (Decimal(trade["calculation_amount_e8"]) / SCALE * quantity / total_e8).quantize(
                            Decimal(".01")
                        )
                    )
                    for _, quantity in followers
                ]
                print(
                    f"[WARN] {trade['broker_code']} 成交 {trade['id']} 人员超分 {Decimal(tolerated_excess) / SCALE} 股，小于1股，按流水保留。"
                )
            else:
                allocations = _allocation_amounts(
                    trade["calculation_amount_e8"], [quantity for _, quantity in followers]
                )
            for (name, quantity_e8), amount_e8 in zip(followers, allocations):
                if name not in people:
                    raise ValueError(
                        f"{trade['broker_code']} 成交 {trade['id']} 的人工归属含数据库不存在的人员。"
                    )
                if quantity_e8 <= 0:
                    continue
                expected.append(
                    dict(
                        person_id=people[name],
                        account_id=trade["account_id"],
                        instrument_id=trade["instrument_id"],
                        broker_trade_id=trade["id"],
                        operation_kind="MARKET_TRADE",
                        side=trade["side"],
                        business_date=trade["business_date"],
                        book_time=trade["book_time"],
                        quantity_e8=quantity_e8,
                        price_e8=trade["price_e8"],
                        amount_e8=amount_e8,
                        amount_basis="LEGACY",
                        currency="USD",
                        link_status="LINKED",
                        legacy_order_id=trade["broker_order_id"] or "",
                        legacy_direction="买入" if trade["side"] == "BUY" else "卖出",
                        legacy_transaction_type="市场交易",
                        legacy_counterparty="",
                    )
                )
            if sum(item["quantity_e8"] for item in expected) != total_e8 + tolerated_excess:
                raise ValueError(f"{trade['broker_code']} 成交 {trade['id']} 的人员份额不守恒，已停止。")
            if (
                not tolerated_excess
                and sum(item["amount_e8"] for item in expected) != trade["calculation_amount_e8"]
            ):
                raise ValueError(f"{trade['broker_code']} 成交 {trade['id']} 的人员金额不守恒，已停止。")
        remaining = {row["id"] for row in existing}
        for values in expected:
            matched = next(
                (
                    row
                    for row in existing
                    if row["id"] in remaining and row["person_id"] == values["person_id"]
                ),
                None,
            )
            if matched is not None:
                remaining.remove(matched["id"])
                if all(matched[key] == value for key, value in values.items()):
                    continue
                fields = list(values)
                connection.execute(
                    f"UPDATE person_operations SET {','.join(key + '=?' for key in fields)},source_record_id=? WHERE id=?",
                    [values[key] for key in fields] + [trade["source_record_id"], matched["id"]],
                )
            else:
                # 恢复曾撤销的同一成交时允许重用逻辑键，不更换券商成交 ID。
                operation_key = f"DB:TRADE:{trade['id']}:PERSON:{values['person_id']}"
                collision = connection.execute(
                    "SELECT id FROM person_operations WHERE operation_key=?", (operation_key,)
                ).fetchone()
                if collision:
                    raise ValueError("数据库人员操作逻辑键冲突，请核对人工编辑。")
                _insert(
                    connection,
                    "person_operations",
                    id=next_id,
                    operation_key=operation_key,
                    sequence_no=next_sequence,
                    source_record_id=trade["source_record_id"],
                    legacy_source="数据库券商成交分摊",
                    **values,
                )
                next_sequence += 1
                next_id += 1
            changed += 1
        changed += _archive_operations(connection, remaining, "交易流水.xlsx修订或撤销人员份额")
    return changed


def _import_trade_flow(connection, base):
    """完整读取已落盘流水；新增、修订、撤销及【谁】均从该 Excel 入库。"""
    from portfolio.sync.sync_and_split_operations import _ticker, _direction

    ledger = TradeFlowWorkbook(base)
    path = ledger.path
    relative = path.relative_to(base).as_posix()
    count = 0
    for broker, sheet in TRADE_SHEETS.items():
        frame = ledger.read(broker)
        scope = f"TRADE_FLOW:{broker}"
        last = connection.execute(
            "SELECT id,content_sha256,parser_version FROM import_batches "
            "WHERE source_scope=? AND source_path=? AND status='SUCCESS' ORDER BY id DESC LIMIT 1",
            (scope, relative),
        ).fetchone()
        if (
            last is not None
            and last["content_sha256"] == ledger.digest
            and last["parser_version"] == PARSER_VERSION
        ):
            ids = {
                row["source_row"]: row["id"]
                for row in connection.execute(
                    "SELECT id,source_row FROM raw_records WHERE import_batch_id=?", (last["id"],)
                )
            }
        else:
            _, ids = _batch(
                connection,
                scope=scope,
                source_kind="EXCEL",
                source_path=relative,
                digest=ledger.digest,
                sheet=sheet,
                records=frame.to_dict("records"),
            )
        normalized = []
        order_ids = set()
        for ordinal, item in enumerate(frame.to_dict("records"), 2):
            if _skip_unsupported_security(connection, ids[ordinal], item.get("Ticker")):
                continue
            identity = normalize_order_id(item.get("订单ID")) or ""
            if identity and identity in order_ids:
                raise ValueError(f"交易流水.xlsx [{sheet}] 第 {ordinal} 行订单 ID 重复，不能重复入账。")
            if identity:
                order_ids.add(identity)
            record = {
                **item,
                "订单ID": identity,
                "__source_row": ordinal,
                "Ticker": _ticker(item.get("Ticker")),
                "方向": _direction(item.get("方向")),
            }
            try:
                when = pd.Timestamp(record["交易时间"])
                numbers = [float(record[key]) for key in ("成交数量", "成交均价", "成交金额")]
                valid = (
                    not pd.isna(when)
                    and when.tzinfo is None
                    and all(math.isfinite(v) for v in numbers)
                    and numbers[0] > 0
                    and numbers[1] >= 0
                    and numbers[2] >= 0
                    and record["方向"] in {"买入", "卖出"}
                )
            except (ValueError, TypeError, KeyError):
                valid = False
            if not valid:
                raise ValueError(f"交易流水.xlsx [{sheet}] 第 {ordinal} 行时间、数量、价格、金额或方向无效。")
            if broker in {"LONGBRIDGE", "SCHWAB"} and not identity and when >= pd.Timestamp(START):
                raise ValueError(f"交易流水.xlsx [{sheet}] 第 {ordinal} 行缺少订单 ID。")
            normalized.append(record)
        desired = pd.DataFrame(normalized, columns=[*TRADE_COLUMNS, "__source_row"])
        count += _apply_trades(
            connection, broker, desired, source_ids=ids, reverse_missing=True, authoritative_source=True
        )
    for sheet, frame in ledger.cash_frames.items():
        broker = "FIRSTRADE" if sheet == "firstrade_资金来源" else "SCHWAB"
        required = (
            {"日期", "交易类别", "金额"}
            if broker == "FIRSTRADE"
            else {"活动ID", "日期", "净金额(USD)", "类型", "说明"}
        )
        if not required.issubset(frame.columns):
            raise ValueError(f"交易流水.xlsx [{sheet}] 缺少字段：{sorted(required - set(frame.columns))}")
        scope = f"TRADE_FLOW:CASH:{broker}"
        last = connection.execute(
            "SELECT id,content_sha256,parser_version FROM import_batches "
            "WHERE source_scope=? AND source_path=? AND status='SUCCESS' ORDER BY id DESC LIMIT 1",
            (scope, relative),
        ).fetchone()
        if (
            last is not None
            and last["content_sha256"] == ledger.digest
            and last["parser_version"] == PARSER_VERSION
        ):
            ids = {
                row["source_row"]: row["id"]
                for row in connection.execute(
                    "SELECT id,source_row FROM raw_records WHERE import_batch_id=?", (last["id"],)
                )
            }
        else:
            _, ids = _batch(
                connection,
                scope=scope,
                source_kind="EXCEL",
                source_path=relative,
                digest=ledger.digest,
                sheet=sheet,
                records=frame.to_dict("records"),
            )
        if broker == "FIRSTRADE":
            _firstrade_cash(
                connection,
                frame,
                _account(connection, broker),
                ids,
                "原始数据/手动维护/交易订单数据firstrade.csv",
            )
        else:
            _apply_schwab_cash(connection, frame, ids)
    count += _import_transfers(connection, base, ledger)
    if hashlib.sha256(path.read_bytes()).hexdigest() != ledger.digest:
        raise ValueError("交易流水.xlsx 在入库期间发生变化，数据库已回滚，请保存并关闭文件后重试。")
    print(
        f"[SYNC] 已从交易流水.xlsx 完整读取成交、资金、内部划转及【谁】，数据库新增、修订或撤销 {count} 笔。"
    )


def _import_transfers(connection, base, ledger):
    """同步交易流水.xlsx中的 [内部划转] 工作表到 person_operations。"""
    frame = getattr(ledger, "transfer_frame", None)
    if frame is None:
        return 0
    from portfolio.sync.sync_and_split_operations import _ticker

    if not frame.empty:
        required = {"交易时间", "渠道", "Ticker", "方向", "成交数量", "划转对象"}
        if not required.issubset(frame.columns):
            raise ValueError(
                f"交易流水.xlsx [内部划转] 缺少必要字段：{sorted(required - set(frame.columns))}"
            )
    scope = "TRADE_FLOW:TRANSFER"
    sheet = "内部划转"
    path = ledger.path
    relative = path.relative_to(base).as_posix()
    last = connection.execute(
        "SELECT id,content_sha256,parser_version FROM import_batches "
        "WHERE source_scope=? AND source_path=? AND status='SUCCESS' ORDER BY id DESC LIMIT 1",
        (scope, relative),
    ).fetchone()
    if (
        last is not None
        and last["content_sha256"] == ledger.digest
        and last["parser_version"] == PARSER_VERSION
    ):
        ids = {
            row["source_row"]: row["id"]
            for row in connection.execute(
                "SELECT id,source_row FROM raw_records WHERE import_batch_id=?", (last["id"],)
            )
        }
    else:
        _, ids = _batch(
            connection,
            scope=scope,
            source_kind="EXCEL",
            source_path=relative,
            digest=ledger.digest,
            sheet=sheet,
            records=frame.to_dict("records"),
        )

    people = {row["name"]: row["id"] for row in connection.execute("SELECT id, name FROM persons")}
    if "个人" not in people:
        raise ValueError("数据库缺少个人账户归属，无法导入内部划转。")
    personal_id = people["个人"]

    next_sequence = connection.execute(
        "SELECT coalesce(max(sequence_no),-1)+1 FROM person_operations"
    ).fetchone()[0]
    next_id = connection.execute("SELECT coalesce(max(id),0)+1 FROM person_operations").fetchone()[0]

    existing_ops = connection.execute(
        "SELECT o.*, r.source_row FROM person_operations o "
        "LEFT JOIN raw_records r ON r.id = o.source_record_id "
        "WHERE o.operation_kind = 'OWNER_TRANSFER' ORDER BY o.id"
    ).fetchall()

    groups = {}
    for op in existing_ops:
        tgk = op["transfer_group_key"]
        if not tgk:
            continue
        if tgk not in groups:
            s_row = op["source_row"]
            if not s_row:
                m = re.search(r":TRANSFER:(\d+)$", tgk)
                if m:
                    s_row = int(m.group(1))
            groups[tgk] = {"source_row": s_row, "ops": {}}
        groups[tgk]["ops"][op["side"]] = op

    unmatched_tgks = set(groups.keys())
    changed = 0

    for idx, item in enumerate(frame.to_dict("records")):
        ordinal = idx + 2
        raw_id = ids[ordinal]
        when = pd.Timestamp(item["交易时间"])
        if pd.isna(when):
            raise ValueError(f"交易流水.xlsx [内部划转] 第 {ordinal} 行交易时间无效。")
        book_time = when.strftime("%Y-%m-%d %H:%M:%S")
        business_date = when.strftime("%Y-%m-%d")
        broker = _text(item["渠道"]).upper()
        account_id = _account(connection, broker)
        symbol = _ticker(item["Ticker"])
        instrument_id = _instrument(connection, symbol, raw_id)
        direction = _text(item["方向"])
        if direction not in ("他人→个人", "个人→他人"):
            raise ValueError(f"交易流水.xlsx [内部划转] 第 {ordinal} 行方向无效：{direction}")
        other_name = _text(item["划转对象"])
        if other_name not in people:
            raise ValueError(f"交易流水.xlsx [内部划转] 第 {ordinal} 行划转对象不存在：{other_name}")
        other_id = people[other_name]

        quantity_val = item.get("成交数量")
        price_val = (
            item.get("成交均价($)")
            if "成交均价($)" in item and not pd.isna(item.get("成交均价($)"))
            else item.get("成交均价")
        )
        amount_val = (
            item.get("成交金额($)")
            if "成交金额($)" in item and not pd.isna(item.get("成交金额($)"))
            else item.get("成交金额")
        )
        if pd.isna(quantity_val) or pd.isna(price_val) or pd.isna(amount_val):
            raise ValueError(f"交易流水.xlsx [内部划转] 第 {ordinal} 行数量、单价或金额无效。")
        quantity_e8 = _scaled(quantity_val)
        price_e8 = _scaled(price_val)
        amount_e8 = _scaled(amount_val)
        tx_type = _text(item.get("交易类型") or "内部划转")

        if direction == "他人→个人":
            out_person, out_cp, out_legacy_cp = other_id, personal_id, "个人"
            in_person, in_cp, in_legacy_cp = personal_id, other_id, other_name
        else:
            out_person, out_cp, out_legacy_cp = personal_id, other_id, other_name
            in_person, in_cp, in_legacy_cp = other_id, personal_id, "个人"

        # 优先匹配同一账户、证券、双方、时间与数量完全相同的记录
        matched_tgk = None
        for tgk in unmatched_tgks:
            g = groups[tgk]
            out_op = g["ops"].get("TRANSFER_OUT")
            if not out_op:
                continue
            if (
                out_op["account_id"] == account_id
                and out_op["instrument_id"] == instrument_id
                and out_op["person_id"] == out_person
                and out_op["counterparty_person_id"] == out_cp
                and out_op["book_time"] == book_time
                and out_op["quantity_e8"] == quantity_e8
            ):
                matched_tgk = tgk
                break

        # 其次匹配同一源行记录（如用户在 Excel 中修订了渠道账户、日期、数量或金额）
        if not matched_tgk:
            for tgk in unmatched_tgks:
                g = groups[tgk]
                if g["source_row"] == ordinal or tgk == f"TRADE_FLOW:TRANSFER:{ordinal}":
                    matched_tgk = tgk
                    break

        transfer_key = matched_tgk if matched_tgk else f"TRADE_FLOW:TRANSFER:{ordinal}"
        if matched_tgk:
            unmatched_tgks.remove(matched_tgk)

        out_values = dict(
            person_id=out_person,
            account_id=account_id,
            instrument_id=instrument_id,
            broker_trade_id=None,
            operation_kind="OWNER_TRANSFER",
            side="TRANSFER_OUT",
            business_date=business_date,
            book_time=book_time,
            executed_at_utc=None,
            quantity_e8=quantity_e8,
            price_e8=price_e8,
            amount_e8=amount_e8,
            amount_basis="TRANSFER_VALUE",
            currency="USD",
            counterparty_person_id=out_cp,
            transfer_group_key=transfer_key,
            link_status="LINKED",
            source_record_id=raw_id,
            legacy_order_id=None,
            legacy_direction=direction,
            legacy_transaction_type=tx_type,
            legacy_source="Money3交易流水唯一交易来源",
            legacy_counterparty=out_legacy_cp,
        )
        in_values = dict(
            person_id=in_person,
            account_id=account_id,
            instrument_id=instrument_id,
            broker_trade_id=None,
            operation_kind="OWNER_TRANSFER",
            side="TRANSFER_IN",
            business_date=business_date,
            book_time=book_time,
            executed_at_utc=None,
            quantity_e8=quantity_e8,
            price_e8=price_e8,
            amount_e8=amount_e8,
            amount_basis="TRANSFER_VALUE",
            currency="USD",
            counterparty_person_id=in_cp,
            transfer_group_key=transfer_key,
            link_status="LINKED",
            source_record_id=raw_id,
            legacy_order_id=None,
            legacy_direction=direction,
            legacy_transaction_type=tx_type,
            legacy_source="Money3交易流水唯一交易来源",
            legacy_counterparty=in_legacy_cp,
        )

        existing_group = groups.get(matched_tgk) if matched_tgk else None
        for side, values in (("TRANSFER_OUT", out_values), ("TRANSFER_IN", in_values)):
            op = existing_group["ops"].get(side) if existing_group else None
            if op is not None:
                needs_update = any(op[k] != v for k, v in values.items() if k in op.keys())
                if needs_update or op["source_record_id"] != raw_id:
                    fields = list(values.keys())
                    connection.execute(
                        f"UPDATE person_operations SET {','.join(f + '=?' for f in fields)} WHERE id=?",
                        [values[f] for f in fields] + [op["id"]],
                    )
                    changed += 1
            else:
                op_key = f"{transfer_key}:{side}"
                _insert(
                    connection,
                    "person_operations",
                    id=next_id,
                    operation_key=op_key,
                    sequence_no=next_sequence,
                    **values,
                )
                next_sequence += 1
                next_id += 1
                changed += 1

    archived_ids = []
    for tgk in unmatched_tgks:
        for side, op in groups[tgk]["ops"].items():
            archived_ids.append(op["id"])
    if archived_ids:
        changed += _archive_operations(connection, archived_ids, "交易流水.xlsx内部划转删除或撤销")

    return changed


def sync_database_orders(base_dir=None, database=None, *, refresh_sources=True):
    from portfolio.sync.database_store import connect_database, database_path, validate_database

    base = Path(base_dir) if base_dir is not None else BASE_DIR
    target = Path(database) if database is not None else database_path()
    validate_database(path=target)
    ledger = TradeFlowWorkbook(base)
    connection = connect_database(target)
    saved = False
    try:
        connection.execute("BEGIN IMMEDIATE")
        batch_start = connection.execute("SELECT coalesce(max(id),0) FROM import_batches").fetchone()[0]
        for broker, sheet in TRADE_SHEETS.items():
            if sheet not in ledger.sheet_names and not _read_trades(connection, broker).empty:
                raise ValueError(
                    f"交易流水.xlsx 缺少已有成交渠道 [{sheet}]，请恢复该工作表后入库，避免误清空。"
                )
        if refresh_sources:
            _import_bbae(connection, base, ledger)
            _import_firstrade(connection, base, ledger)
            _sync_longbridge(connection, ledger)
            _sync_schwab(connection, ledger)
            # 先保存流水，再从磁盘重新读取；券商适配器不直接写 broker_trades。
            ledger.save()
        saved = bool(ledger.changed or ledger.cash_changed or ledger.cancelled_changed)
        _import_trade_flow(connection, base)
        count = refresh_allocations(connection)
        from portfolio.sync.trade_fees import refresh_fees
        refresh_fees(connection)
        if connection.execute("PRAGMA foreign_key_check").fetchone() is not None:
            raise ValueError("同步后数据库外键核对未通过，本次修改已回滚。")
        connection.execute(
            "UPDATE import_batches SET status='SUCCESS' WHERE status='PENDING' AND parser_version=? AND id>?",
            (PARSER_VERSION, batch_start),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        if saved:
            print("[ERROR] 数据库未提交；已保存的交易流水.xlsx 保留，请修正问题后重新运行入库。")
        raise
    finally:
        connection.close()
    print(f"[DONE] 交易流水 → 数据库同步完成，人员成交份额新增、修订或归档 {count} 笔；独立人工划转已保留。")
    return [target]

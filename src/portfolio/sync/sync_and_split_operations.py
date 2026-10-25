"""Money3本地同步组件；从既有账本适配器迁入并裁剪非同步入口。"""

from __future__ import annotations
import re


import pandas as pd


import openpyxl


from openpyxl.styles import Font, PatternFill


from openpyxl.utils import get_column_letter


from portfolio.sync.order_id_utils import normalize_order_id


SPLIT_START = pd.Timestamp("2026-07-01 00:00:00")


TICKER_ALIASES = {"SKHYV": "SKHY"}


def _number(value: object) -> float:
    if pd.isna(value):
        return 0.0
    text = str(value).strip().replace("$", "").replace(",", "")
    if text.startswith("(") and text.endswith(")"):
        text = f"-{text[1:-1]}"
    number = pd.to_numeric(text, errors="coerce")
    return 0.0 if pd.isna(number) else float(number)


def _ticker(value: object) -> str:
    ticker = str(value or "").strip().upper().split(".")[0].replace(".0", "")
    return TICKER_ALIASES.get(ticker, ticker)


def _to_eastern_naive(value: object, source_timezone: str = "Asia/Shanghai") -> pd.Timestamp:
    timestamp = pd.to_datetime(value, errors="coerce")
    if pd.isna(timestamp):
        return pd.NaT
    if timestamp.tzinfo is None:
        timestamp = timestamp.tz_localize(source_timezone)
    return timestamp.tz_convert("America/New_York").tz_localize(None)


def _beijing_bookkeeping_time(value: object) -> pd.Timestamp:
    """将以北京时间记录的订单转换为本账本的交易归属时间。

    账务日从北京时间 08:00 开始、至次日 07:59:59 结束；因此整体前移 8 小时。
    仅用于确定所属交易日，并非美东时钟换算。
    """
    timestamp = pd.to_datetime(value, errors="coerce")
    if pd.isna(timestamp):
        return pd.NaT
    if timestamp.tzinfo is not None:
        timestamp = timestamp.tz_convert("Asia/Shanghai").tz_localize(None)
    return timestamp - pd.Timedelta(hours=8)


def _direction(value: object) -> str:
    text = str(value).strip()
    if any(token in text.lower() for token in ("buy", "bought")) or "买" in text:
        return "买入"
    if any(token in text.lower() for token in ("sell", "sold")) or "卖" in text:
        return "卖出"
    return text


def _refresh_firstrade_source_window(
    existing: pd.DataFrame,
    downloaded: pd.DataFrame,
    source_window: tuple[pd.Timestamp, pd.Timestamp] | None = None,
) -> tuple[pd.DataFrame, int, int, int, pd.Timestamp, pd.Timestamp]:
    """Refresh the date window covered by the Firstrade CSV.

    Firstrade exports may correct historical fills.  Keep rows before/after the
    current export window, but replace the covered window from source while
    carrying forward any manually assigned owner (``谁``) on matched orders.
    """
    base_columns = ["交易时间", "Ticker", "方向", "成交数量", "成交均价", "成交金额", "谁"]
    existing = existing.copy()
    downloaded = downloaded.copy()
    for frame in (existing, downloaded):
        if "谁" not in frame.columns:
            frame["谁"] = ""
        frame["交易时间"] = pd.to_datetime(frame.get("交易时间"), errors="coerce")

    downloaded = downloaded.dropna(subset=["交易时间"]).copy()
    if downloaded.empty and source_window is None:
        return existing, 0, 0, 0, pd.NaT, pd.NaT
    if (downloaded["交易时间"] < SPLIT_START).any():
        raise ValueError("第一证券导入区间早于增量起点，禁止覆盖历史底稿")

    # Identical rows can be separate executions, not duplicate imports.
    # Idempotency comes from replacing the complete source window, not deduping fills.
    def fill_key(row):
        return (
            pd.Timestamp(row["交易时间"]).isoformat(),
            _ticker(row.get("Ticker")),
            str(row.get("方向", "")).strip(),
            format(abs(_number(row.get("成交数量"))), ".10g"),
            format(_number(row.get("成交均价")), ".10g"),
            f"{abs(_number(row.get('成交金额'))):.2f}",
        )

    if source_window is None:
        window_start = downloaded["交易时间"].min().normalize()
        window_end = downloaded["交易时间"].max().normalize()
    else:
        window_start, window_end = (pd.Timestamp(value).normalize() for value in source_window)
        if (
            pd.isna(window_start)
            or pd.isna(window_end)
            or window_start < SPLIT_START
            or window_end < window_start
        ):
            raise ValueError("第一证券导入区间无效，禁止覆盖历史底稿")
    historical = existing[(existing["交易时间"] < SPLIT_START) | existing["交易时间"].isna()].copy()
    incremental = existing[existing["交易时间"] >= SPLIT_START].copy()
    in_window = incremental[
        (incremental["交易时间"] >= window_start)
        & (incremental["交易时间"] < window_end + pd.Timedelta(days=1))
    ].copy()
    retained = incremental.drop(index=in_window.index)

    owners_by_key = {}
    for _, row in in_window.iterrows():
        value = row.get("谁", "")
        owner = "" if pd.isna(value) else str(value).strip()
        owners_by_key.setdefault(fill_key(row), []).append(owner)
    seen = {}
    owners = []
    for _, row in downloaded.iterrows():
        key = fill_key(row)
        old_owners = owners_by_key.get(key, [])
        if len({owner for owner in old_owners if owner}) > 1:
            raise ValueError(
                f"第一证券同值成交存在多个归属，无法唯一匹配：{key}；本次第一证券同步未写入，请核对成交及归属"
            )
        occurrence = seen.get(key, 0)
        seen[key] = occurrence + 1
        value = row.get("谁", "")
        incoming_owner = "" if pd.isna(value) else str(value).strip()
        old_owner = old_owners[occurrence] if occurrence < len(old_owners) else ""
        if old_owner and incoming_owner and old_owner != incoming_owner:
            raise ValueError(f"第一证券成交归属冲突：{key}，本次同步未写入")
        owners.append(old_owner or incoming_owner)
    # Do not silently discard manually assigned executions on a reduced/partial export.
    for key, old_owners in owners_by_key.items():
        if any(old_owners[seen.get(key, 0) :]):
            raise ValueError(
                f"第一证券导出缺少已有归属的成交：{key}；"
                "可能为部分导出或成交修订，请确认完整记录及归属后再同步"
            )
    downloaded["谁"] = owners
    merged = pd.concat([historical, retained, downloaded], ignore_index=True, sort=False)
    merged = merged.reindex(columns=[*base_columns, *[c for c in merged.columns if c not in base_columns]])
    merged = merged.sort_values("交易时间", ascending=False, kind="stable").reset_index(drop=True)
    return merged, len(downloaded), len(in_window), len(retained), window_start, window_end


def _bbae_order_identity(row: pd.Series, time_field: str = "交易时间", date_only: bool = False) -> str:
    """BBAE 无订单 ID 时用于识别同一笔订单；成交价也必须参与防止误合并。"""
    timestamp = pd.to_datetime(row.get(time_field), errors="coerce")
    if pd.isna(timestamp):
        time_text = ""
    elif date_only:
        time_text = timestamp.strftime("%Y-%m-%d")
    else:
        time_text = timestamp.strftime("%Y-%m-%d %H:%M:%S")
    return "|".join(
        [
            time_text,
            _ticker(row.get("Ticker")),
            str(row.get("方向", "")).strip(),
            f"{_number(row.get('成交数量')):.4f}",
            f"{_number(row.get('成交均价')):.4f}",
        ]
    )


def _preserve_bbae_who(
    old_rows: pd.DataFrame, source_rows: pd.DataFrame
) -> tuple[pd.DataFrame, int, list[dict[str, object]]]:
    """BBAE 重同步时保留人工“谁”；多候选归属不猜测，留空并返回待确认清单。"""
    old = old_rows.copy()
    source = source_rows.copy()
    if "谁" not in old.columns:
        old["谁"] = ""
    if "交易时间" not in old.columns:
        old["交易时间"] = pd.NaT
    old["交易时间"] = pd.to_datetime(old["交易时间"], errors="coerce")

    exact_who: dict[str, set[str]] = {}
    date_who: dict[str, set[str]] = {}
    for _, row in old.iterrows():
        who = str(row.get("谁", "")).strip()
        if not who or who.lower() == "nan":
            continue
        exact_who.setdefault(_bbae_order_identity(row), set()).add(who)
        date_who.setdefault(_bbae_order_identity(row, date_only=True), set()).add(who)

    preserved = 0
    ambiguous: list[dict[str, object]] = []
    for index, row in source.iterrows():
        candidates: set[str] = set()
        # 兼容已按账务日写入过的时间，以及旧版直接写入的原始北京时间。
        candidates.update(exact_who.get(_bbae_order_identity(row), set()))
        candidates.update(exact_who.get(_bbae_order_identity(row, time_field="__原始北京时间"), set()))
        # 曾按其他时区写入的旧数据：仅在订单特征及账务日期唯一时作为兜底。
        candidates.update(date_who.get(_bbae_order_identity(row, date_only=True), set()))
        if len(candidates) == 1:
            source.at[index, "谁"] = next(iter(candidates))
            preserved += 1
        elif len(candidates) > 1:
            source.at[index, "谁"] = ""
            ambiguous.append(
                {
                    "交易时间": row["交易时间"],
                    "Ticker": row["Ticker"],
                    "方向": row["方向"],
                    "成交数量": row["成交数量"],
                    "候选谁": "；".join(sorted(candidates)),
                }
            )
    return source, preserved, ambiguous


def _reconcile_bbae_orders(
    old_rows: pd.DataFrame,
    source_rows: pd.DataFrame,
) -> tuple[pd.DataFrame, int, list[dict[str, object]], list[dict[str, object]]]:
    """仅校正能安全匹配的 BBAE 旧行，保留原始文件未覆盖的既有流水。"""
    old = old_rows.copy()
    if "谁" not in old.columns:
        old["谁"] = ""
    if "交易时间" not in old.columns:
        old["交易时间"] = pd.NaT
    old["交易时间"] = pd.to_datetime(old["交易时间"], errors="coerce")
    source, preserved_count, ambiguous_who = _preserve_bbae_who(old, source_rows)

    exact_indices: dict[str, set[int]] = {}
    date_indices: dict[str, set[int]] = {}
    for index, row in old.iterrows():
        exact_indices.setdefault(_bbae_order_identity(row), set()).add(index)
        date_indices.setdefault(_bbae_order_identity(row, date_only=True), set()).add(index)

    matched_old_indices: set[int] = set()
    accepted_source_indices: list[int] = []
    unresolved: list[dict[str, object]] = []
    for index, row in source.iterrows():
        candidates = set()
        candidates.update(exact_indices.get(_bbae_order_identity(row), set()))
        candidates.update(exact_indices.get(_bbae_order_identity(row, time_field="__原始北京时间"), set()))
        candidates.update(exact_indices.get(_bbae_order_identity(row, time_field="__旧版美东时间"), set()))
        if not candidates:
            date_candidates = date_indices.get(_bbae_order_identity(row, date_only=True), set())
            if len(date_candidates) == 1:
                candidates.update(date_candidates)
            elif len(date_candidates) > 1:
                # 同一天同订单特征出现多笔旧行时，不能安全替换任意一笔。
                unresolved.append(
                    {
                        "交易时间": row["交易时间"],
                        "Ticker": row["Ticker"],
                        "方向": row["方向"],
                        "成交数量": row["成交数量"],
                        "原因": "同日存在多笔候选旧订单",
                    }
                )
                continue
        matched_old_indices.update(candidates)
        accepted_source_indices.append(index)

    retained_old = old.drop(index=list(matched_old_indices))
    accepted_source = source.loc[accepted_source_indices].drop(
        columns=["__原始北京时间", "__旧版美东时间"], errors="ignore"
    )
    merged = pd.concat([retained_old, accepted_source], ignore_index=True, sort=False)
    if not merged.empty:
        merged["交易时间"] = pd.to_datetime(merged["交易时间"], errors="coerce")
        merged = merged.sort_values("交易时间", ascending=False).reset_index(drop=True)
    return merged, preserved_count, ambiguous_who, unresolved


def is_ignored_ticker(ticker: str) -> bool:
    if str(ticker or "").strip().upper().endswith(".HK"):
        return True
    t = _ticker(ticker)
    return not t or t.isdigit() or t.endswith(".HK")


def _longbridge_context(*, strict: bool = False):
    try:
        from portfolio.sync.longbridge_credentials import load_longbridge_credentials

        LONGBRIDGE_ACCESS_TOKEN, LONGBRIDGE_APP_KEY, LONGBRIDGE_APP_SECRET = load_longbridge_credentials()
    except (OSError, ValueError) as exc:
        if strict:
            raise RuntimeError(
                f"无法读取 Money3 长桥凭证：{exc}。请修复配置后重试；"
                "如只使用数据库现有账本，可添加 --skip-sync。"
            ) from None
        print(f"[WARN] 无法读取 Money3 长桥凭证: {exc}")
        return None
    try:
        from longbridge.openapi import Config, TradeContext

        config = Config.from_apikey(
            app_key=LONGBRIDGE_APP_KEY,
            app_secret=LONGBRIDGE_APP_SECRET,
            access_token=LONGBRIDGE_ACCESS_TOKEN,
        )
        return TradeContext(config)
    except Exception as exc:
        if strict:
            # SDK 异常可能带连接地址，避免把完整异常串中的凭证写入日志。
            code = re.search(r"\bcode\s*[=:]\s*(\d+)", str(exc))
            detail = f"错误码 {code.group(1)}" if code else type(exc).__name__
            reason = (
                "，Token 已过期，请更新长桥凭证"
                if "token expired" in str(exc).lower()
                else "，请检查凭证、网络及 SDK 安装"
            )
            raise RuntimeError(
                f"长桥交易连接初始化失败（{detail}）{reason}。本次数据库同步未提交；"
                "如只使用数据库现有账本，可添加 --skip-sync。"
            ) from None
        return None


def _write_df_to_sheet(wb: openpyxl.Workbook, sheet_name: str, df: pd.DataFrame) -> None:
    df = df.copy()
    if "订单ID" in df.columns:
        df["订单ID"] = df["订单ID"].map(normalize_order_id)
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(sheet_name)

    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            h = str(ws.cell(1, c).value or "")
            if h in ("成交均价", "成交金额"):
                cell.number_format = "$#,##0.00;[Red]-$#,##0.00;-"
            elif h == "成交数量":
                cell.number_format = "#,##0.00;[Red]-#,##0.00;-"
            elif h == "交易时间":
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif h == "订单ID":
                cell.value = normalize_order_id(cell.value)
                cell.number_format = "@"

    for col in range(1, ws.max_column + 1):
        col_let = get_column_letter(col)
        ws.column_dimensions[col_let].width = 16


def _same_transaction_rows(left: pd.DataFrame, right: pd.DataFrame) -> bool:
    """比较交易流水内容，忽略行顺序和 Excel/Pandas 的日期、数值类型差异。"""
    columns = ["交易时间", "Ticker", "方向", "成交数量", "成交均价", "成交金额", "谁", "订单ID"]

    def canonical(frame: pd.DataFrame) -> pd.DataFrame:
        result = frame.reindex(columns=columns).copy()
        result["交易时间"] = (
            pd.to_datetime(result["交易时间"], errors="coerce").dt.strftime("%Y-%m-%d %H:%M:%S").fillna("")
        )
        for column in ("成交数量", "成交均价", "成交金额"):
            result[column] = pd.to_numeric(result[column], errors="coerce").fillna(0.0).astype(float).round(8)
        for column in ("Ticker", "方向", "谁"):
            result[column] = (
                result[column].fillna("").astype(str).str.strip().replace({"nan": "", "None": ""})
            )
        result["订单ID"] = result["订单ID"].map(normalize_order_id).fillna("")
        return result.sort_values(columns, kind="stable").reset_index(drop=True)

    return canonical(left).equals(canonical(right))
